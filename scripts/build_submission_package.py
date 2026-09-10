import csv
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission_ready"
FIG_OUT = OUT / "figures"
CSV_OUT = OUT / "csv"


TITLE = "Cooperative Co Evolution Hybrid Optimization for Dynamic Task Scheduling in Heterogeneous Distributed Systems"
AUTHORS = "Divyanshu Kashyap, Om Yadav, and Tejas Joshi"
AFFILIATION = "Netaji Subhas University of Technology, Delhi, India"


def reset_output():
    OUT.mkdir(exist_ok=True)
    FIG_OUT.mkdir(exist_ok=True)
    CSV_OUT.mkdir(exist_ok=True)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text_color(cell, color="FFFFFF"):
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(color)


def set_table_borders(table, color="D9D9D9"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(table, top=80, start=80, bottom=80, end=80):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for m, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = margins.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)

    for name, size in [
        ("Title", 17),
        ("Heading 1", 14),
        ("Heading 2", 12),
        ("Heading 3", 11),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)


def add_para(doc, text="", style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.05
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        lead.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(8 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table)
    set_cell_margins(table)
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = str(header)
        set_cell_shading(hdr[i], "1F4E79")
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for paragraph in hdr[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(font_size)
                run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cells[i], "F4F8FB")
            for paragraph in cells[i].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table


def add_figure(doc, image_path, caption, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(9)


def fmt(value, digits=4):
    if isinstance(value, str):
        return value
    try:
        return f"{float(value):.{digits}f}"
    except Exception:
        return value


def rows_from_df(df, columns, digits=4):
    rows = []
    for _, row in df.iterrows():
        rows.append([fmt(row[col], digits) for col in columns])
    return rows


def build_manuscript():
    full = pd.read_csv(ROOT / "results/csv/phase2_full_comparison.csv").rename(columns={"Unnamed: 0": "Scheduler"})
    summary = pd.read_csv(ROOT / "results/csv/final_summary_table.csv")
    weights = pd.read_csv(ROOT / "results/csv/phase2_best_weights.csv")
    sensitivity = pd.read_csv(ROOT / "results/csv/fitness_weight_sensitivity.csv")
    robustness = pd.read_csv(ROOT / "results/csv/robustness_scenarios.csv")

    doc = Document()
    style_doc(doc)
    title = doc.add_paragraph(TITLE, style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = add_para(doc, AUTHORS)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = add_para(doc, AFFILIATION)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "Abstract", 1)
    add_para(
        doc,
        "Objectives: This study addresses dynamic task scheduling in heterogeneous distributed systems where task arrivals, queue states, node speeds, communication latency, and failures change over time. Methods and analysis: A policy based scheduler is developed in which six nonnegative weights guide node selection across processing speed, queue workload, network latency, failure probability, available capacity, and execution cost. The weights are optimized using Genetic Algorithm, Particle Swarm Optimization, Differential Evolution, and a Cooperative Co Evolution Hybrid optimizer. Performance is evaluated against Round Robin, Least Loaded, and MECT baselines using average completion time, P95 latency, makespan, load variance, failure rate, throughput, and a revised Schedule Length Ratio based on a parallel lower bound. Findings: On the primary 1000 task Alibaba trace sample, optimized policies reduce average completion time by 14.59 to 21.27 percent and P95 latency by 4.06 to 13.79 percent relative to MECT. Differential Evolution provides the best P95, makespan, SLR, and throughput, while Particle Swarm Optimization provides the best average completion time and failure rate. Novelty and improvement: The work contributes a transparent six objective scheduling formulation, constrained penalty weights that preserve feature semantics, a corrected SLR definition, and robustness testing across samples, arrival windows, workload intensities, and cluster sizes."
    )
    add_para(doc, "Keywords: Task scheduling; heterogeneous distributed systems; cooperative co evolution; metaheuristics; Alibaba Cluster Trace; multi objective optimization")

    add_heading(doc, "1 Introduction", 1)
    add_para(doc, "Dynamic distributed systems execute workloads whose arrivals, resource needs, queue states, and failure conditions change continuously. Heterogeneous clusters intensify this problem because fast, low latency nodes can reduce completion time but may also concentrate load, while slower or less reliable nodes can damage tail performance. Classical heuristics such as Round Robin and Least Loaded remain attractive because they are simple and inexpensive, but they cannot fully represent the competing objectives that govern modern cloud and edge scheduling.")
    add_para(doc, "Metaheuristic scheduling methods, including Genetic Algorithm, Particle Swarm Optimization, Differential Evolution, and hybrid approaches, are widely used for search spaces where exact optimization is impractical. However, many studies still rely on static workflow assumptions, synthetic workloads, or narrow objective functions. Recent work in edge and cloud scheduling has increasingly emphasized deep reinforcement learning, SLA awareness, energy and cost trade offs, and real time resource optimization, but the need remains for interpretable scheduling policies that can be tested against trace derived dynamic arrivals [19-24].")
    add_para(doc, "This paper revises and strengthens the original scheduling framework in three ways. First, the policy search is formulated as a six objective normalized fitness function rather than a four term proxy. Second, all policy weights are constrained to be nonnegative so that penalty features cannot reverse meaning during optimization. Third, SLR is redefined against a parallel lower bound, making it meaningful for independent tasks that execute concurrently. The resulting evaluation is intentionally presented as a trade off analysis rather than a universal dominance claim.")
    add_para(doc, "The main contributions are: (1) an interpretable six feature policy scheduler for heterogeneous nodes; (2) a corrected six objective fitness function that includes AvgCT, P95, makespan, load variance, failure rate, and SLR; (3) a Cooperative Co Evolution Hybrid optimizer that decomposes the policy vector into throughput, overhead, and reliability subspaces; (4) real Alibaba Cluster Trace 2018 evaluation with reproducible data download and checksum verification; and (5) sensitivity and robustness analyses requested during peer review.")

    add_heading(doc, "2 Related Work", 1)
    add_para(doc, "Task scheduling in heterogeneous systems has a long history. HEFT and related list scheduling heuristics provide effective low complexity baselines for workflow scheduling [1-3]. Foundational distributed systems and cloud computing work established the importance of heterogeneity, resource allocation, and scalability [4, 5], while scheduling complexity results explain why large scheduling instances motivate heuristic and metaheuristic approaches [6].")
    add_para(doc, "Population based metaheuristics such as GA [7, 8], PSO [9], and DE [11] provide flexible search mechanisms for difficult optimization landscapes. Hybrid approaches, including GA PSO and EDA GA variants, can improve exploration and convergence [14, 15]. Cooperative co evolution decomposes a high dimensional search problem into interacting subcomponents and is therefore well suited to weight vectors with functionally distinct groups [16].")
    add_para(doc, "Recent work has expanded task scheduling research toward edge and cloud continuum environments. Aminu et al. reviewed metaheuristic task scheduling strategies in edge computing [19]. Naeem et al. studied real time job scheduling and resource optimization in edge cloud continuum networks [20]. Zhang and Ou presented reinforcement learning for energy efficient and cost effective cloud edge scheduling [21]. Sravan and Shaik proposed DRL based multi objective scheduling for latency, energy, and SLA objectives [22]. Yamsani and Reddy examined SLA aware deep reinforcement learning for adaptive edge cloud task scheduling [23], while Raju et al. proposed an IntelliScheduler framework based on hybrid deep learning [24]. These studies motivate the revised paper's emphasis on dynamic behavior, multiple objectives, and robustness across workload settings.")

    add_heading(doc, "3 Methodology", 1)
    add_heading(doc, "3.1 Dataset And Workload Construction", 2)
    add_para(doc, "The experiments use the official Alibaba Cluster Trace 2018 batch task file. The downloaded archive was verified with SHA256 value 7c4b32361bd1ec2083647a8f52a6854a03bc125ca5c202652316c499fbf978c6. The raw file contains 14,295,731 task rows. After retaining terminated records and removing invalid start time or CPU entries, 14,058,462 records remain. A 1000 task sample is selected for the main experiment, arrivals are rescaled to a 600 second window, and planned CPU is log normalized into workload units.")
    add_table(doc, ["Property", "Value"], [
        ["Source dataset", "Alibaba Cluster Trace 2018 batch_task.csv"],
        ["Raw records", "14,295,731"],
        ["Cleaned records", "14,058,462"],
        ["Tasks sampled", "1,000"],
        ["Arrival window", "0.0 to 600.0 seconds"],
        ["Workload range", "10.0 to 84.4 units"],
        ["Workload mean", "54.72 units"],
        ["Archive SHA256", "7c4b32361bd1ec2083647a8f52a6854a03bc125ca5c202652316c499fbf978c6"],
    ], widths=[2.0, 4.6], font_size=8.2)

    add_heading(doc, "3.2 Simulation Environment", 2)
    add_para(doc, "The simulated cluster contains 15 heterogeneous nodes distributed across high, medium, and low performance tiers. Node speeds, communication latencies, and failure probabilities are sampled from tier specific ranges using fixed random seeds. The default tier structure is five high performance nodes, seven medium performance nodes, and three low performance nodes.")
    add_table(doc, ["Tier", "Nodes", "Speed range", "Latency range", "Failure probability"], [
        ["High", "5", "8 to 12 units/s", "5 to 20 ms", "0.00 to 0.01"],
        ["Medium", "7", "4 to 8 units/s", "15 to 40 ms", "0.01 to 0.03"],
        ["Low", "3", "2 to 4 units/s", "30 to 60 ms", "0.02 to 0.05"],
    ], widths=[1.0, 0.8, 1.5, 1.5, 1.7])
    add_para(doc, "Queue workload is updated continuously as Q_j(t) = max(0, Q_j(t-1) - delta_t s_j), where s_j is the processing speed of node j. Completion time is CT_ij = latency_j/1000 + workload_i/s_j + Q_j/s_j. Failed tasks are retried up to two times with a short delay and are then assigned to the least loaded node if all retries fail.")

    add_heading(doc, "3.3 Policy Scheduler", 2)
    add_para(doc, "For each arriving task, the policy scheduler scores every node using six normalized features. The selected node is the node with the highest score.")
    add_para(doc, "Score_j = w1 speed_j - w2 queue_j - w3 latency_j - w4 failure_j + w5 capacity_j - w6 cost_j")
    add_para(doc, "The optimizer searches W = [w1, w2, w3, w4, w5, w6] within [0, 5]^6. Nonnegative weights preserve feature semantics: speed and capacity are rewards, while queue, latency, failure probability, and execution cost remain penalties.")
    add_table(doc, ["Weight", "Feature", "Role"], [
        ["w1", "Processing speed", "Reward"],
        ["w2", "Queue workload", "Penalty"],
        ["w3", "Network latency", "Penalty"],
        ["w4", "Failure probability", "Penalty"],
        ["w5", "Available capacity", "Reward"],
        ["w6", "Execution cost", "Penalty"],
    ], widths=[0.8, 2.5, 1.2])

    add_heading(doc, "3.4 Notation And Parameters", 2)
    add_para(doc, "The following table defines the symbols and parameters used in the scheduling model, simulator, fitness function, and robustness tests.")
    add_table(doc, ["Symbol", "Definition"], [
        ["i", "Task index"],
        ["j", "Node index"],
        ["w1 to w6", "Policy weights for speed, queue workload, latency, failure probability, capacity, and cost"],
        ["s_j", "Processing speed of node j"],
        ["Q_j(t)", "Queue workload assigned to node j at time t"],
        ["L_j", "Communication latency of node j in milliseconds"],
        ["p_j", "Failure probability of node j"],
        ["C_j", "Execution cost feature of node j"],
        ["A_i", "Arrival time of task i"],
        ["X_i", "Workload of task i"],
        ["CT_ij", "Completion time if task i is assigned to node j"],
        ["P95", "Ninety fifth percentile task completion latency"],
        ["M", "Observed makespan from first task arrival to final task finish"],
        ["LB_parallel", "Parallel lower bound used for SLR"],
        ["SLR", "Schedule Length Ratio, calculated as M divided by LB_parallel"],
        ["r_k", "Metric k normalized by the MECT value from the same workload and node realization"],
        ["F", "Overall six objective fitness value minimized by the optimizers"],
        ["N", "Number of tasks in an experiment"],
        ["T_window", "Arrival window used to rescale sampled trace arrivals"],
        ["alpha", "Workload scale multiplier used in robustness testing"],
    ], widths=[1.3, 5.3], font_size=7.8)

    add_heading(doc, "4 Fitness Function", 1)
    add_para(doc, "The final fitness function optimizes all six stated objectives. Each objective is normalized by the MECT baseline from the same workload and node realization.")
    add_para(doc, "F = 0.25 r_AvgCT + 0.20 r_P95 + 0.20 r_Makespan + 0.10 r_LoadVariance + 0.15 r_FailureRate + 0.10 r_SLR")
    add_para(doc, "A soft failure guard adds 0.05(r_FailureRate - 1) when the normalized failure rate exceeds 1.0. This structure gives priority to mean latency, tail latency, and makespan while retaining reliability, load balance, and schedule length efficiency.")
    add_para(doc, "SLR is computed as observed makespan divided by a parallel lower bound. The lower bound is the maximum of the aggregate work capacity bound, the release time bound based on fastest feasible task durations, and the longest single task bound. This avoids the invalid serial denominator used in the previous manuscript.")
    add_table(doc, ["Objective", "Coefficient", "Meaning"], [
        ["Average completion time", "0.25", "Primary mean latency objective"],
        ["P95 latency", "0.20", "Tail latency objective"],
        ["Makespan", "0.20", "Total schedule length"],
        ["Load variance", "0.10", "Load balance pressure"],
        ["Failure rate", "0.15", "Reliability objective"],
        ["SLR", "0.10", "Normalized schedule efficiency"],
    ], widths=[2.4, 1.1, 3.0])

    add_heading(doc, "5 Optimizers", 1)
    add_para(doc, "GA, PSO, and DE search the complete six dimensional policy vector. GA uses tournament selection, uniform crossover, Gaussian mutation, and elitism. PSO uses inertia, cognitive, and social terms with bounded velocities. DE uses the DE/rand/1 strategy with binomial crossover.")
    add_para(doc, "The Cooperative Co Evolution Hybrid decomposes the vector into three two dimensional groups: throughput weights (w1, w5) assigned to DE, overhead weights (w2, w3) assigned to PSO, and reliability weights (w4, w6) assigned to GA. A shared global best vector coordinates the sub optimizers, and the global best is injected every ten iterations.")
    add_table(doc, ["Parameter", "GA", "PSO", "DE", "Hybrid CCE"], [
        ["Population or swarm", "20", "20", "20", "20 per subgroup"],
        ["Iterations", "50", "50", "50", "50"],
        ["Crossover rate", "0.8", "-", "0.9", "GA 0.8, DE 0.9"],
        ["Mutation", "0.1", "-", "F = 0.5", "GA 0.15, DE F = 0.5"],
        ["PSO parameters", "-", "w = 0.7, c1 = 1.5, c2 = 1.5", "-", "same"],
    ], widths=[1.6, 1.0, 1.7, 1.0, 1.8], font_size=7.8)

    add_heading(doc, "6 Experimental Setup", 1)
    add_para(doc, "The experiments were implemented in Python using NumPy, Pandas, and Matplotlib. The main result uses 1000 tasks sampled from the official Alibaba trace and a 15 node heterogeneous cluster. Random seeds are fixed for reproducibility. Baseline constants are generated from MECT and passed directly to the optimizer through a shared metrics object.")
    add_para(doc, "The evaluation reports average completion time, P95 latency, load variance, failure rate, throughput, makespan, SLR lower bound, and SLR. Additional analyses include coefficient sensitivity and robustness across repeated samples, 300, 600, and 1200 second arrival windows, light and heavy workload scales, and 10 and 25 node clusters.")

    add_heading(doc, "7 Results", 1)
    add_heading(doc, "7.1 Baseline Performance", 2)
    baseline = full[full["Scheduler"].isin(["Round Robin", "Least Loaded", "MECT"])]
    add_table(doc, ["Scheduler", "AvgCT", "P95", "LoadVar", "FailRate", "Throughput", "Makespan", "SLR"], rows_from_df(baseline, ["Scheduler", "avg_completion_time", "p95_latency", "load_variance", "failure_rate", "throughput", "makespan", "slr"]), font_size=7.6)
    add_para(doc, "Round Robin produces the lowest load variance because it distributes tasks cyclically without regard to node state, but it performs poorly on completion time, P95, makespan, throughput, and SLR. MECT provides the strongest heuristic baseline for tail latency and schedule length, while Least Loaded reduces average completion time but has worse P95 and makespan than MECT.")
    add_para(doc, "Figure 1 visualizes these baseline trade offs and shows why MECT is used as the normalization reference for the optimizer fitness values.")
    add_figure(doc, ROOT / "results/plots/baseline_results.png", "Figure 1. Baseline heuristic scheduler comparison.", width=6.2)

    add_heading(doc, "7.2 Optimizer Performance", 2)
    add_table(doc, ["Scheduler", "AvgCT", "P95", "LoadVar", "FailRate", "Throughput", "Makespan", "SLR"], rows_from_df(full, ["Scheduler", "avg_completion_time", "p95_latency", "load_variance", "failure_rate", "throughput", "makespan", "slr"]), font_size=7.0)
    add_para(doc, "On the main Alibaba trace sample, optimized policies reduce average completion time by 14.59 to 21.27 percent and P95 latency by 4.06 to 13.79 percent relative to MECT. PSO provides the best average completion time at 9.4771 s and ties GA for the best failure rate at 0.010. DE provides the best P95 latency at 16.2014 s, best makespan at 617.5510 s, best SLR at 1.0210, and best throughput at 1.6193.")
    add_para(doc, "The results do not support a universal dominance claim. Hybrid CCE remains competitive on mean and tail latency, but it regresses on failure rate, makespan, SLR, and throughput compared with MECT in this run. This is a key revision from the earlier manuscript and is reported transparently.")
    add_para(doc, "Figure 2 compares the main metrics across all schedulers, while Figure 3 reports convergence and Figure 4 shows the subgroup behavior inside the Cooperative Co Evolution Hybrid optimizer.")
    add_figure(doc, ROOT / "results/plots/scheduler_comparison.png", "Figure 2. Scheduler performance comparison across AvgCT, P95, makespan, and failure rate.", width=6.2)
    add_figure(doc, ROOT / "results/plots/convergence_curves.png", "Figure 3. Convergence curves for all optimizers.", width=5.9)
    add_figure(doc, ROOT / "results/plots/cce_subgroup_convergence.png", "Figure 4. Cooperative Co Evolution subgroup convergence.", width=5.9)
    add_para(doc, "Figure 5 summarizes percentage improvement against MECT and makes the metric specific gains and regressions explicit.")
    add_figure(doc, ROOT / "results/plots/improvement_over_mect.png", "Figure 5. Percentage improvement over MECT. Positive values indicate improvement.", width=6.1)

    add_heading(doc, "7.3 Weight Interpretation", 2)
    add_table(doc, ["Optimizer", "Fitness", "w1", "w2", "w3", "w4", "w5", "w6"], rows_from_df(weights, ["optimizer", "fitness", "w1", "w2", "w3", "w4", "w5", "w6"]), font_size=7.8)
    add_para(doc, "All optimized policies assign high importance to the queue penalty w2. PSO, DE, and CCE drive w2 to the upper bound of 5.0, and GA reaches 4.4855. The available capacity reward w5 also receives substantial weight in all methods. This supports the interpretation that queue avoidance and capacity preference are central for dynamic trace based scheduling.")
    add_para(doc, "Figure 6 displays the learned nonnegative weights and confirms that no penalty feature receives a sign reversing coefficient.")
    add_figure(doc, ROOT / "results/plots/weight_vectors_heatmap.png", "Figure 6. Learned nonnegative policy weights for all optimizers.", width=5.8)

    add_heading(doc, "8 Sensitivity And Robustness", 1)
    sens_cols = ["profile", "avg_completion_time", "p95_latency", "makespan", "slr", "load_variance", "failure_rate"]
    add_table(doc, ["Profile", "AvgCT", "P95", "Makespan", "SLR", "LoadVar", "FailRate"], rows_from_df(sensitivity, sens_cols), font_size=6.6)
    add_para(doc, "The coefficient sensitivity study indicates that changing the weights can improve one target at the expense of another. The fairness heavy setting lowers load variance but worsens P95, makespan, SLR, and failure rate. Tail and makespan profiles improve specific objectives but do not dominate the balanced six objective configuration. The balanced profile is therefore retained because it avoids the strongest single metric regressions.")
    rob_rows = []
    for scenario, group in robustness.groupby("scenario", sort=False):
        mect = group[group["scheduler"] == "MECT"].iloc[0]
        de = group[group["scheduler"] == "Reoptimized DE Policy"].iloc[0]
        rob_rows.append([
            scenario,
            f"{(mect.avg_completion_time - de.avg_completion_time) / mect.avg_completion_time * 100:.2f}%",
            f"{(mect.p95_latency - de.p95_latency) / mect.p95_latency * 100:.2f}%",
            f"{(mect.makespan - de.makespan) / mect.makespan * 100:.2f}%",
            f"{(mect.slr - de.slr) / mect.slr * 100:.2f}%",
            f"{(mect.failure_rate - de.failure_rate) / max(mect.failure_rate, 1e-9) * 100:.2f}%",
        ])
    add_table(doc, ["Scenario", "AvgCT", "P95", "Makespan", "SLR", "FailRate"], rob_rows, widths=[1.7, 0.9, 0.9, 1.0, 0.9, 1.0], font_size=6.8)
    add_para(doc, "Robustness tests show a nuanced result. Reoptimized DE improves average completion time by 6.66 percent on average across nine scenarios and improves failure rate by 12.07 percent on average, but it regresses on makespan and SLR by 2.16 percent on average. The 300 second compressed arrival window, 1200 second relaxed window, and 10 node cluster reveal sensitivity to workload dynamics and cluster scale.")
    add_para(doc, "Figure 7 isolates the revised SLR values, Figure 8 gives a multi metric radar profile, and Figure 9 relates optimizer runtime to final fitness.")
    add_figure(doc, ROOT / "results/plots/slr_comparison.png", "Figure 7. Revised SLR comparison using the parallel lower bound.", width=5.4)
    add_figure(doc, ROOT / "results/plots/radar_profile.png", "Figure 8. Multi metric radar profile for MECT and optimized schedulers.", width=5.4)
    add_figure(doc, ROOT / "results/plots/runtime_vs_fitness.png", "Figure 9. Runtime versus best fitness for optimizer efficiency comparison.", width=5.4)

    add_heading(doc, "9 Comparison With Prior Work", 1)
    add_para(doc, "The closest prior comparison in the original manuscript was Li and Chen [18], who studied energy aware task scheduling using a metaheuristic approach. A direct numerical SLR comparison is not retained because their workflow based SLR formulation is not methodologically equivalent to the parallel lower bound based SLR used here. Instead, the revised manuscript compares the proposed method qualitatively with prior metaheuristic and reinforcement learning studies.")
    add_para(doc, "Compared with recent edge cloud and cloud edge scheduling work [19-24], this study emphasizes interpretability and reviewer reproducibility: the scheduler uses an explicit six feature score, the data are fetched from the official Alibaba trace archive, and the robustness experiments expose settings where the method regresses. This strengthens the credibility of the result even though it narrows the claim.")

    add_heading(doc, "10 Conclusion", 1)
    add_para(doc, "This study presented a dynamic scheduling framework for heterogeneous distributed systems in which task assignment is controlled by an interpretable six feature policy and optimized through evolutionary and cooperative co evolutionary search. The revised formulation resolves the mismatch between the original six objective claim and the previous four term fitness function by explicitly incorporating average completion time, P95 latency, makespan, load variance, failure rate, and Schedule Length Ratio into a single normalized objective. The resulting policy search remains easy to interpret because each learned weight corresponds to a concrete scheduling feature, while the nonnegative weight bounds preserve the intended meaning of penalty terms for queue length, latency, failure probability, and execution cost. The corrected SLR metric also avoids an invalid serial denominator by comparing observed schedule length against a parallel lower bound derived from task releases, fastest feasible task durations, and aggregate cluster capacity.")
    add_para(doc, "The experimental workflow now supports a stronger evaluation protocol. Baseline constants are generated once from the MECT run and reused consistently across the optimizer fitness calculation, result tables, and manuscript text. Sensitivity analysis justifies the selected coefficient values, and robustness experiments vary random samples, arrival compression windows, workload intensity, and cluster size to test whether the conclusions survive changes in workload dynamics. On the primary real trace setting, optimized policies substantially improve mean and tail latency relative to MECT. However, robustness tests show that these improvements are not uniform across all operating conditions, particularly under compressed arrivals and smaller clusters. Future work should improve robustness under heavier contention, evaluate larger samples, compare against directly comparable workflow benchmarks, add energy and cost measurements where data are available, and validate the policy in a real distributed testbed.")

    add_heading(doc, "References", 1)
    refs = [
        "Topcuoglu H, Hariri S, Wu MY (2002) Performance effective and low complexity task scheduling for heterogeneous computing. IEEE Transactions on Parallel and Distributed Systems 13(3):260-274. https://doi.org/10.1109/71.993206",
        "Ibrahim OH, Ahmed I, Omara FA, Shaker FA (2008) An enhanced min min algorithm for job scheduling in grid computing. Proceedings of the 5th International Conference on Electrical Engineering, pp 1-9.",
        "Mao Y, Chen X, Li X (2014) Max Min task scheduling algorithm for load balance in cloud computing. Proceedings of Computer Science and Network Technology, pp 457-463. https://doi.org/10.1007/978-81-322-1759-6_53",
        "Tanenbaum AS, van Steen M (2007) Distributed Systems Principles and Paradigms. Pearson Prentice Hall.",
        "Buyya R, Yeo CS, Venugopal S, Broberg J, Brandic I (2009) Cloud computing and emerging IT platforms. Future Generation Computer Systems 25(6):599-616. https://doi.org/10.1016/j.future.2008.12.001",
        "Ullman JD (1975) NP complete scheduling problems. Journal of Computer and System Sciences 10(3):384-393. https://doi.org/10.1016/S0022-0000(75)80008-0",
        "Holland JH (1975) Adaptation in Natural and Artificial Systems. University of Michigan Press.",
        "Goldberg DE (1989) Genetic Algorithms in Search Optimization and Machine Learning. Addison Wesley.",
        "Kennedy J, Eberhart R (1995) Particle swarm optimization. Proceedings of IEEE International Conference on Neural Networks, pp 1942-1948. https://doi.org/10.1109/ICNN.1995.488968",
        "Wang Y, Zuo X (2021) An effective cloud workflow scheduling approach combining PSO and idle time slot aware rules. IEEE CAA Journal of Automatica Sinica 8(5):1079-1094. https://doi.org/10.1109/JAS.2021.1003982",
        "Storn R, Price K (1997) Differential evolution a simple and efficient heuristic for global optimization over continuous spaces. Journal of Global Optimization 11(4):341-359. https://doi.org/10.1023/A:1008202821328",
        "Deb K, Pratap A, Agarwal S, Meyarivan T (2002) A fast and elitist multiobjective genetic algorithm NSGA II. IEEE Transactions on Evolutionary Computation 6(2):182-197. https://doi.org/10.1109/4235.996017",
        "Durillo L, Fard HM, Prodan R (2012) MOHEFT a multi objective list based method for workflow scheduling. Proceedings of IEEE CloudCom, pp 185-192. https://doi.org/10.1109/CloudCom.2012.6427573",
        "Senthil Kumar AM, Parthiban K, Siva Shankar S (2019) An efficient task scheduling in a cloud computing environment using hybrid Genetic Algorithm Particle Swarm Optimization algorithm. Proceedings of ICISS, IEEE, pp 1-6. https://doi.org/10.1109/ISS1.2019.8908041",
        "Pang S, Li W, He H, Shan Z, Wang X (2019) An EDA GA hybrid algorithm for multi objective task scheduling in cloud computing. IEEE Access 7:146379-146389. https://doi.org/10.1109/ACCESS.2019.2946216",
        "Potter MA, De Jong KA (1994) A cooperative coevolutionary approach to function optimization. Proceedings of PPSN III, pp 249-257. https://doi.org/10.1007/3-540-58484-6_269",
        "Alibaba Cluster Trace Program (2018) Alibaba Cluster Data V2018. Alibaba Group. https://github.com/alibaba/clusterdata",
        "Li C, Chen L (2024) Optimization for energy aware design of task scheduling in heterogeneous distributed systems a meta heuristic based approach. Computing. https://doi.org/10.1007/s00607-024-01282-1",
        "Aminu J, Latip R, Hanafi ZM, Kamarudin S, Gabi D (2025) Systematic review of metaheuristic based task scheduling strategies in edge computing environments. Discover Computing 28:307. https://doi.org/10.1007/s10791-025-09819-4",
        "Naeem AB, Senapati B, Rasheed J, Baili J, Osman O (2025) An intelligent job scheduling and real time resource optimization for edge cloud continuum in next generation networks. Scientific Reports 15:41534. https://doi.org/10.1038/s41598-025-25452-z",
        "Zhang W, Ou H (2025) Reinforcement learning based multi objective task scheduling for energy efficient and cost effective cloud edge computing. Scientific Reports 15:41716. https://doi.org/10.1038/s41598-025-25666-1",
        "Sravan P, Shaik MA (2026) DRL based multi objective task scheduling for edge cloud computing latency energy and SLA optimisation. Scientific Reports 16:19681. https://doi.org/10.1038/s41598-026-49824-1",
        "Yamsani N, Reddy CP (2026) SLA aware deep reinforcement learning for adaptive EdgeCloud task scheduling. Scientific Reports 16:10037. https://doi.org/10.1038/s41598-026-40237-8",
        "Raju LR, Reddy MVK, Surukanti SR, Sudhakar G, Sarma MVVS, Adepu A (2026) IntelliScheduler an edge cloud computing environment hybrid deep learning framework for task scheduling based on learning. Scientific Reports 16:11219. https://doi.org/10.1038/s41598-026-41330-8",
    ]
    for i, ref in enumerate(refs, start=1):
        add_para(doc, f"[{i}] {ref}")

    out = OUT / "Revised_Manuscript.docx"
    doc.save(out)
    return out


def build_response():
    doc = Document()
    style_doc(doc)
    title = doc.add_paragraph("Response To Reviewers", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, "Manuscript title: " + TITLE)
    add_para(doc, "Authors: " + AUTHORS)
    add_para(doc, "Dear Editor and Reviewers,")
    add_para(doc, "We thank the reviewers and technical editor for their careful reading and constructive comments. The manuscript, experiments, and code have been revised to address the methodological and presentation concerns. The major changes include a six objective fitness function, constrained nonnegative scheduling weights, a corrected SLR definition, consistent MECT normalization constants, real Alibaba trace reruns, coefficient sensitivity analysis, robustness experiments, updated figures and tables, expanded result interpretation, and revised manuscript text.")
    sections = [
        ("Reviewer 1", [
            ("English revisions", "The manuscript has been edited throughout for grammar, clarity, and technical style. If the journal requires a language certificate, the authors should upload the certificate from their chosen editing service."),
            ("Abstract format", "The abstract has been revised as a single structured paragraph covering Objectives, Methods and analysis, Findings, and Novelty and improvement."),
            ("Introduction and literature gap", "The introduction now explains the gap around dynamic heterogeneous scheduling, trace derived arrivals, and robustness validation."),
            ("Recent work", "Six recent 2025-2026 references have been added to the related work and reference list."),
            ("Figure quality", "Figures have been regenerated from the experiment outputs and included in the revised Word manuscript."),
            ("Results explanation", "The results now discuss trade offs instead of claiming universal dominance. Improvements and regressions are reported transparently."),
            ("Comparison with previous studies", "The invalid direct SLR comparison with Li and Chen has been removed and replaced with a qualitative comparison."),
            ("Conclusion", "The conclusion has been rewritten to summarize the corrected objective, findings, limitations, robustness results, and future work."),
            ("Figures and symbols", "All figures and charts are cited in the manuscript text, and a notation table has been added."),
            ("References and format", "References have been numbered sequentially and the manuscript has been prepared as a Word file."),
        ]),
        ("Reviewer 3", [
            ("Six objectives", "The revised fitness function includes AvgCT, P95, makespan, load variance, failure rate, and SLR directly."),
            ("Coefficient rationale", "A coefficient sensitivity analysis was added and run on the real Alibaba trace sample."),
            ("MECT normalization", "MECT metrics are generated once and reused consistently through set_baselines_from_metrics."),
            ("SLR definition", "SLR now uses a parallel lower bound rather than a serial sum of fastest task times."),
            ("Li and Chen comparison", "The direct numerical SLR comparison has been removed because the definitions are not methodologically comparable."),
            ("Negative weights", "The weight search range is now [0, 5], preserving the intended penalty semantics."),
            ("Robustness", "Robustness tests cover repeated samples, 300/600/1200 second windows, light/heavy workloads, and 10/15/25 node clusters."),
        ]),
        ("Technical Editor", [
            ("Word version", "A revised Word manuscript is included in the submission package."),
            ("ORCID", "At least one author ORCID must be entered by the submitting author in the portal. This field cannot be invented by the assistant."),
            ("Language editing", "The revised manuscript has been rewritten and edited throughout for clarity. The authors should check the portal proof and attach a language editing certificate if the journal requires one."),
        ]),
    ]
    for heading, items in sections:
        add_heading(doc, heading, 1)
        for comment, response in items:
            add_para(doc, "Comment: " + comment, bold_lead="Comment:")
            add_para(doc, "Response: " + response, bold_lead="Response:")
    out = OUT / "Response_to_Reviewers.docx"
    doc.save(out)
    return out


def build_cover_letter():
    doc = Document()
    style_doc(doc)
    title = doc.add_paragraph("Cover Letter For Revised Manuscript", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, "Dear Editor,")
    add_para(doc, "Please find enclosed the revised manuscript titled \"" + TITLE + "\" by " + AUTHORS + ". We appreciate the detailed comments from the reviewers and technical editor. The revision substantially strengthens the paper's methodology, reproducibility, and presentation.")
    add_para(doc, "The main revisions are: the fitness function now includes all six stated objectives; scheduling weights are constrained to nonnegative values; SLR has been redefined using a parallel lower bound; MECT normalization constants are generated consistently; the experiments have been rerun on the official Alibaba Cluster Trace 2018 batch task file; coefficient sensitivity and robustness analyses have been added; and the results have been rewritten to report both improvements and regressions transparently.")
    add_para(doc, "On the primary 1000 task Alibaba trace sample, optimized policies reduce average completion time by 14.59 to 21.27 percent and P95 latency by 4.06 to 13.79 percent relative to MECT. The revised manuscript avoids the previous invalid numerical SLR comparison with Li and Chen and frames the contribution as an interpretable, trace evaluated adaptive scheduling framework.")
    add_para(doc, "The submission package includes the revised manuscript, response to reviewers, regenerated figures, generated CSV tables, and an author metadata checklist. At least one author ORCID should be entered in the submission system before final upload.")
    add_para(doc, "Sincerely,")
    add_para(doc, "Divyanshu Kashyap, Om Yadav, and Tejas Joshi")
    out = OUT / "Cover_Letter.docx"
    doc.save(out)
    return out


def build_metadata_checklist():
    doc = Document()
    style_doc(doc)
    title = doc.add_paragraph("Author Metadata And Upload Checklist", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, "This document lists the final portal fields and upload materials for the revised submission.")
    add_table(doc, ["Field", "Entry"], [
        ["Title", TITLE],
        ["Authors", AUTHORS],
        ["Affiliation", AFFILIATION],
        ["Corresponding author", "Confirm in portal"],
        ["Author ORCID", "Required before upload: enter at least one verified author ORCID"],
        ["Main article file", "Revised_Manuscript.docx"],
        ["Response file", "Response_to_Reviewers.docx"],
        ["Cover letter", "Cover_Letter.docx"],
        ["Figures", "submission_ready/figures"],
        ["Result CSV files", "submission_ready/csv"],
    ], widths=[2.2, 4.4])
    add_heading(doc, "Upload Steps", 1)
    steps = [
        "Open the revision task in the journal submission portal.",
        "Upload Revised_Manuscript.docx as the main manuscript file.",
        "Upload Response_to_Reviewers.docx as the reviewer response file.",
        "Upload Cover_Letter.docx if the portal requests a cover letter.",
        "Upload figures separately if the portal asks for figure source files.",
        "Enter at least one verified ORCID in the author metadata page.",
        "Review the portal generated PDF proof carefully before final submission.",
    ]
    for step in steps:
        add_para(doc, step)
    out = OUT / "Author_Metadata_and_Upload_Checklist.docx"
    doc.save(out)
    return out


def copy_support_files():
    for image in (ROOT / "results/plots").glob("*.png"):
        shutil.copy2(image, FIG_OUT / image.name)
    for csv_file in (ROOT / "results/csv").glob("*.csv"):
        shutil.copy2(csv_file, CSV_OUT / csv_file.name)
    manifest = OUT / "SUBMISSION_PACKAGE_MANIFEST.txt"
    manifest.write_text(
        "\n".join([
            "Submission package contents",
            "Revised_Manuscript.docx",
            "Response_to_Reviewers.docx",
            "Cover_Letter.docx",
            "Author_Metadata_and_Upload_Checklist.docx",
            "figures/",
            "csv/",
            "",
            "Final portal action still required: enter at least one verified author ORCID.",
        ]),
        encoding="utf-8",
    )


def main():
    reset_output()
    outputs = [
        build_manuscript(),
        build_response(),
        build_cover_letter(),
        build_metadata_checklist(),
    ]
    copy_support_files()
    for path in outputs:
        print(path)


if __name__ == "__main__":
    main()
